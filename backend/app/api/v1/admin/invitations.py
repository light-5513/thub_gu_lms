"""Invitations: admins send email and/or WhatsApp setup links singly or in bulk."""

import asyncio
import csv
import io
import re
from datetime import datetime, timezone
from typing import Any

import bson
from fastapi import APIRouter, Depends, File, Request, UploadFile
from openpyxl import load_workbook
from pydantic import ValidationError

from app.core.deps import get_db, require_admin
from app.core.errors import AppError
from app.core.rate_limit import client_ip, is_rate_limited
from app.models.enums import AuditAction
from app.schemas.auth import BulkInviteRequest, InviteRecipient, InviteRequest
from app.services.audit_service import AuditService
from app.services.invitation_service import InvitationService
from app.services.whatsapp_service import WhatsAppService

router = APIRouter(prefix="/admin/invitations", tags=["admin-invitations"])
MAX_UPLOAD_BYTES = 2 * 1024 * 1024
MAX_UPLOAD_ROWS = 500
EMAIL_HEADERS = {"email", "email address", "student email", "student email address"}
PHONE_HEADERS = {
    "phone", "phone number", "mobile", "mobile number", "whatsapp", "whatsapp number"
}


@router.post("")
async def invite_student(
    payload: InviteRequest,
    request: Request,
    user=Depends(require_admin),
    db=Depends(get_db),
):
    service = InvitationService(db)
    channels = payload.channels
    if "channels" not in payload.model_fields_set:
        channels = ["email", "whatsapp"] if payload.phone else ["email"]
    return await service.invite(
        payload.email,
        user,
        phone=payload.phone,
        role=payload.role,
        ip=client_ip(request),
        channels=channels,
    )


@router.post("/bulk")
async def bulk_invite(
    payload: BulkInviteRequest,
    request: Request,
    user=Depends(require_admin),
    db=Depends(get_db),
):
    """Send invitations to a batch of email addresses and optional phone numbers."""
    limited, retry_after = await is_rate_limited(
        f"bulk-invite:{client_ip(request)}", limit=10, window_seconds=600
    )
    if limited:
        raise AppError(
            f"Too many bulk invites. Try again in {retry_after} seconds.", 429
        )

    service = InvitationService(db)
    results: list[dict] = []
    seen: set = set()

    recipients = payload.recipients or [
        InviteRecipient(email=email, phone=None) for email in (payload.emails or [])
    ]
    channels = payload.channels
    if "channels" not in payload.model_fields_set:
        channels = (
            ["email", "whatsapp"]
            if payload.recipients and any(recipient.phone for recipient in payload.recipients)
            else ["email"]
        )
    for recipient in recipients:
        email = str(recipient.email).lower().strip()
        if not email:
            continue
        if email in seen:
            results.append(
                {"email": email, "ok": False, "message": "Duplicate in your list"}
            )
            continue
        seen.add(email)
        try:
            res = await service.invite(
                email,
                user,
                phone=recipient.phone,
                role=payload.role,
                ip=client_ip(request),
                channels=channels,
            )
            results.append(
                {
                    "email": email,
                    "ok": True,
                    "message": res["message"],
                    "channels": res["channels"],
                    "partial": res["partial"],
                }
            )
        except AppError as exc:
            # Already registered / already pending — skip without aborting the batch.
            asyncio.ensure_future(
                AuditService(db).log(
                    AuditAction.INVITATION_SENT,
                    user_id=str(user["_id"]),
                    role=user["role"],
                    entity="invitation",
                    entity_id=email,
                    new_data={"email": email, "skipped": exc.message},
                )
            )
            results.append({"email": email, "ok": False, "message": exc.message})
        except Exception:
            results.append({"email": email, "ok": False, "message": "Failed to send"})

    accepted = sum(1 for r in results if r["ok"])
    return {
        "total": len(results),
        "accepted": accepted,
        "partial": sum(1 for r in results if r.get("partial")),
        "failed": len(results) - accepted,
        "results": results,
    }


@router.post("/bulk/parse")
async def parse_invitation_upload(
    file: UploadFile = File(...),
    user=Depends(require_admin),
):
    """Parse a CSV or XLSX into a reviewable invitation recipient list."""
    filename = (file.filename or "").lower()
    if not filename.endswith((".csv", ".xlsx")):
        raise AppError("Upload a .csv or .xlsx file", 400)

    content = await file.read(MAX_UPLOAD_BYTES + 1)
    if len(content) > MAX_UPLOAD_BYTES:
        raise AppError("The upload must be 2 MB or smaller", 413)
    if not content:
        raise AppError("The uploaded file is empty", 400)

    try:
        rows = _read_invitation_rows(filename, content)
    except AppError:
        raise
    except Exception:
        raise AppError("Could not read this file. Check that it is a valid CSV or XLSX.", 400)

    items: list[dict[str, str | None]] = []
    invalid_rows: list[dict[str, int | str]] = []
    seen: set[str] = set()
    data_rows = 0
    header_checked = False
    email_column = 0
    phone_column: int | None = 1
    for row_number, row in rows:
        if not any(_cell_text(cell) for cell in row):
            continue

        if not header_checked:
            detected_email, detected_phone = _column_indexes(row)
            header_checked = True
            if detected_email is not None:
                email_column = detected_email
                phone_column = detected_phone
                continue

        data_rows += 1
        if data_rows > MAX_UPLOAD_ROWS:
            raise AppError("The upload can contain at most 500 non-empty rows", 400)

        email = _cell_at(row, email_column).lower()
        phone = _cell_at(row, phone_column)
        if not email:
            invalid_rows.append({"row": row_number, "message": "Email address is missing"})
            continue
        try:
            recipient = InviteRecipient(email=email, phone=phone or None)
        except ValidationError:
            invalid_rows.append({"row": row_number, "message": "Enter a valid email address"})
            continue
        if phone and not 8 <= len(re.sub(r"\D", "", phone)) <= 15:
            invalid_rows.append(
                {"row": row_number, "message": "Use an international WhatsApp number with 8 to 15 digits"}
            )
            continue

        clean_email = str(recipient.email).lower()
        if clean_email in seen:
            invalid_rows.append({"row": row_number, "message": "Duplicate email address"})
            continue
        seen.add(clean_email)
        items.append({"email": clean_email, "phone": recipient.phone})

    return {"items": items, "invalid_rows": invalid_rows}


@router.get("/whatsapp/status")
async def whatsapp_status(user=Depends(require_admin)):
    return await WhatsAppService.connection_status()


@router.get("/pending")
async def pending_invitations(user=Depends(require_admin), db=Depends(get_db)):
    """List pending (unused, unexpired) invitations for visibility."""
    cursor = (
        db.invitations.find(
            {
                "used": False,
                "expires_at": {"$gt": datetime.now(timezone.utc).replace(tzinfo=None)},
            }
        )
        .sort([("created_at", -1)])
        .limit(200)
    )
    items: list[Any] = []
    async for doc in cursor:
        items.append(
            {
                "id": str(doc["_id"]),
                "email": doc["email"],
                "created_at": _iso(doc.get("created_at")),
                "expires_at": _iso(doc.get("expires_at")),
                "email_status": doc.get("email_status"),
                "email_error": doc.get("email_error"),
                "whatsapp_status": doc.get("whatsapp_status"),
                "whatsapp_error": doc.get("whatsapp_error"),
            }
        )
    return {"items": items}


@router.delete("/{invitation_id}")
async def revoke_invitation(
    invitation_id: str, user=Depends(require_admin), db=Depends(get_db)
):
    try:
        result = await db.invitations.delete_one({"_id": bson.ObjectId(invitation_id)})
    except Exception:
        raise AppError("Invalid invitation id", 400)
    if result.deleted_count == 0:
        raise AppError("Invitation not found", 404)
    return {"message": "Invitation revoked"}


def _iso(dt) -> Any:
    if isinstance(dt, datetime):
        return dt.isoformat()
    return None


def _read_invitation_rows(filename: str, content: bytes):
    if filename.endswith(".csv"):
        text = content.decode("utf-8-sig")
        sample = text[:4096]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(io.StringIO(text, newline=""), dialect)
        return list(enumerate(reader, start=1))

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        return list(
            (row_number, list(values))
            for row_number, values in enumerate(
                sheet.iter_rows(max_row=MAX_UPLOAD_ROWS + 2, values_only=True),
                start=1,
            )
        )
    finally:
        workbook.close()


def _column_indexes(row) -> tuple[int | None, int | None]:
    normalized = [_cell_text(value).lower().replace("_", " ") for value in row]
    email_index = next((i for i, value in enumerate(normalized) if value in EMAIL_HEADERS), None)
    if email_index is None:
        return None, None
    phone_index = next((i for i, value in enumerate(normalized) if value in PHONE_HEADERS), None)
    return email_index, phone_index


def _cell_at(row, index: int | None) -> str:
    if index is None or index >= len(row):
        return ""
    return _cell_text(row[index])


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value)).strip()
    return str(value).strip()

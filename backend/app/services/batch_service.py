"""Batch service for managing student batches and bulk assignments."""

import logging
from io import BytesIO

from bson import ObjectId
from openpyxl import load_workbook
from pymongo import UpdateOne

from app.core.errors import AppError, ConflictError, NotFoundError
from app.models.enums import AuditAction
from app.repositories.batches import BatchRepository
from app.repositories.users import StudentRepository
from app.services.audit_service import AuditService

logger = logging.getLogger(__name__)

class BatchService:
    def __init__(self, db):
        self.db = db
        self.batches = BatchRepository(db)
        self.students = db.students
        self.audit = AuditService(db)

    async def list_batches(self, page: int = 1, page_size: int = 50) -> dict:
        """List all batches."""
        total = await self.batches.collection.count_documents({})
        skip = (page - 1) * page_size
        docs = await self.batches.collection.find({}).sort("created_at", -1).skip(skip).limit(page_size).to_list(length=page_size)
        batch_ids = [str(doc["_id"]) for doc in docs]
        counts: dict[str, int] = {}
        if batch_ids:
            rows = await self.students.aggregate([
                {"$match": {"batch_id": {"$in": batch_ids}}},
                {"$group": {"_id": "$batch_id", "count": {"$sum": 1}}},
            ]).to_list(length=page_size)
            counts = {str(row["_id"]): row["count"] for row in rows}
        items = []
        from datetime import datetime
        for doc in docs:
            doc["id"] = str(doc.pop("_id"))
            doc["student_count"] = counts.get(doc["id"], 0)
            for k, v in doc.items():
                if isinstance(v, datetime):
                    doc[k] = v.isoformat()
            items.append(doc)
            
        return {
            "items": items,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": (total + page_size - 1) // page_size
        }

    async def create_batch(self, batch_data: dict) -> dict:
        """Create a new batch."""
        return await self.batches.create_batch(batch_data)

    async def get_batch(self, batch_id: str) -> dict:
        """Get a batch by ID."""
        if not ObjectId.is_valid(batch_id):
            raise AppError("Invalid batch ID format", 400)
        batch = await self.batches.get(batch_id)
        batch["student_count"] = await self.students.count_documents({"batch_id": batch_id})
        return batch

    async def delete_batch(self, batch_id: str) -> None:
        """Delete a batch and unassign all students."""
        await self.get_batch(batch_id)
        if await self.db.classes.count_documents({"batch_id": batch_id}):
            raise AppError("Delete this batch's classes before deleting the batch", 409)
        await self.batches.delete(batch_id)
        # Unassign students
        await self.students.update_many({"batch_id": batch_id}, {"$set": {"batch_id": None}})

    async def assign_students(self, batch_id: str, roll_numbers: list[str]) -> dict:
        """Assign a list of roll numbers to a batch."""
        # Ensure batch exists
        batch = await self.get_batch(batch_id)
        
        # Find valid roll numbers
        cursor = self.students.find({"roll_number": {"$in": roll_numbers}}, {"roll_number": 1, "batch_id": 1})
        found_students = await cursor.to_list(None)
        found_rolls = {s["roll_number"] for s in found_students}
        previous_batch_ids = {
            s.get("batch_id") for s in found_students
            if s.get("batch_id") and s.get("batch_id") != batch_id
        }
        
        failed_rolls = list(set(roll_numbers) - found_rolls)
        success_count = len(found_rolls)
        
        if found_rolls:
            await self.students.update_many(
                {"roll_number": {"$in": list(found_rolls)}},
                {"$set": {"batch_id": batch_id}}
            )
            # Update student count (recalculate completely to be safe)
            actual_count = await self.students.count_documents({"batch_id": batch_id})
            await self.batches.update_batch(batch_id, {"student_count": actual_count})
            for previous_id in previous_batch_ids:
                if isinstance(previous_id, str) and ObjectId.is_valid(previous_id):
                    previous_count = await self.students.count_documents({"batch_id": previous_id})
                    await self.batches.collection.update_one(
                        {"_id": ObjectId(previous_id)},
                        {"$set": {"student_count": previous_count}},
                    )
            
        return {
            "success_count": success_count,
            "failed_roll_numbers": failed_rolls
        }

    async def remove_student(
        self,
        batch_id: str,
        student_id: str,
        actor: dict | None = None,
        ip: str | None = None,
    ) -> None:
        """Unassign one student without deleting their profile or attendance."""
        await self.get_batch(batch_id)
        if not ObjectId.is_valid(student_id):
            raise AppError("Invalid student ID format", 400)
        student_oid = ObjectId(student_id)
        student = await self.students.find_one({"_id": student_oid}, {"batch_id": 1})
        if not student:
            raise NotFoundError("Student not found")
        if student.get("batch_id") != batch_id:
            raise ConflictError("This student is not assigned to this batch")

        result = await self.students.update_one(
            {"_id": student_oid, "batch_id": batch_id},
            {"$set": {"batch_id": None}},
        )
        if not result.modified_count:
            raise ConflictError("This student's batch assignment changed; refresh and try again")

        count = await self.students.count_documents({"batch_id": batch_id})
        await self.batches.collection.update_one(
            {"_id": ObjectId(batch_id)}, {"$set": {"student_count": count}}
        )
        if actor:
            await self.audit.log(
                AuditAction.STUDENT_UPDATED,
                user_id=str(actor["_id"]),
                role=actor["role"],
                entity="student",
                entity_id=student_id,
                old_data={"batch_id": batch_id},
                new_data={"batch_id": None},
                ip_address=ip,
            )

    async def parse_and_assign_xlsx(self, batch_id: str, content: bytes) -> dict:
        """Parse an XLSX file for roll numbers and assign them."""
        wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        sheet = wb.active
        
        # Assume first row is header, look for "roll_number"
        headers = [str(cell.value).strip().lower() for cell in sheet[1] if cell.value]
        if "roll_number" not in headers:
            raise AppError("Missing 'roll_number' column in Excel file")
            
        roll_index = headers.index("roll_number")
        
        roll_numbers = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) > roll_index and row[roll_index]:
                roll_numbers.append(str(row[roll_index]).strip())
                
        if not roll_numbers:
            raise AppError("No roll numbers found in the Excel file")
            
        return await self.assign_students(batch_id, roll_numbers)

    async def get_batch_students(self, batch_id: str, page: int = 1, page_size: int = 50) -> dict:
        """List students in a batch."""
        await self.get_batch(batch_id)
        total = await self.students.count_documents({"batch_id": batch_id})
        skip = (page - 1) * page_size
        cursor = self.students.find({"batch_id": batch_id}).sort("first_name", 1).skip(skip).limit(page_size)
        items = []
        from datetime import datetime
        async for doc in cursor:
            doc["id"] = str(doc.pop("_id"))
            for k, v in doc.items():
                if isinstance(v, datetime):
                    doc[k] = v.isoformat()
            items.append(doc)
            
        return {
            "items": items,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": (total + page_size - 1) // page_size if page_size else 0
        }

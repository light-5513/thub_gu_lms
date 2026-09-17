"""Bulk student import from Excel (.xlsx) with preview/validate/confirm flow."""

from datetime import datetime, timezone
from io import BytesIO
from typing import Any

import bson
from openpyxl import load_workbook

from app.core.errors import AppError, NotFoundError
from app.models.enums import AuditAction
from app.services.audit_service import AuditService
from app.services.student_service import StudentService

REQUIRED_COLUMNS = [
    "first_name",
    "last_name",
    "email",
    "roll_number",
    "course",
    "branch",
    "section",
]
OPTIONAL_COLUMNS = ["phone", "batch", "academic_year"]
MAX_FILE_SIZE = 10 * 1024 * 1024  # 10 MB
MAX_ROWS = 5000


def parse_xlsx(content: bytes) -> list[dict]:
    """Parse an uploaded workbook into a list of row dicts."""
    if len(content) > MAX_FILE_SIZE:
        raise AppError("File is too large (max 10 MB)")
    try:
        wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    except Exception:
        raise AppError("The file could not be parsed as an .xlsx workbook")

    if not rows:
        raise AppError("The workbook has no data")

    header = [str(c).strip().lower() if c else "" for c in rows[0]]
    missing = [c for c in REQUIRED_COLUMNS if c not in header]
    if missing:
        raise AppError(f"Missing required columns: {', '.join(missing)}")

    col_index = {name: i for i, name in enumerate(header)}
    parsed: list[dict] = []
    for i, row in enumerate(rows[1:], start=2):
        record: dict[str, Any] = {}
        for col in set(REQUIRED_COLUMNS + OPTIONAL_COLUMNS):
            if col in col_index and col_index[col] < len(row):
                value = row[col_index[col]]
                if value is not None:
                    value = str(value).strip()
                    if value.lower() in ("", "none", "null"):
                        value = ""
                record[col] = value or None
            else:
                record[col] = None
        record["_row"] = i
        if any(record.get(c) for c in REQUIRED_COLUMNS):
            parsed.append(record)
    return parsed


def validate_rows(
    rows: list[dict], existing_emails: set, existing_rolls: set
) -> tuple[list[dict], list[dict]]:
    """Validate parsed rows; returns (valid, errors). Duplicates within file are flagged."""
    import re

    email_re = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
    valid: list[dict] = []
    errors: list[dict] = []
    seen_emails: set = set()
    seen_rolls: set = set()

    for row in rows:
        r = row["_row"]
        problems: list[tuple[str, str]] = []

        first = (row.get("first_name") or "").strip()
        last = (row.get("last_name") or "").strip()
        email = (row.get("email") or "").lower().strip()
        roll = (row.get("roll_number") or "").strip()
        course = (row.get("course") or "").strip()
        branch = (row.get("branch") or "").strip()
        section = (row.get("section") or "").strip()

        if not first:
            problems.append(("first_name", "Missing First Name"))
        if not last:
            problems.append(("last_name", "Missing Last Name"))
        if not email:
            problems.append(("email", "Missing Email"))
        elif not email_re.match(email):
            problems.append(("email", "Invalid Email"))
        elif email in existing_emails or email in seen_emails:
            problems.append(
                (
                    "email",
                    "Duplicate Email"
                    if email in seen_emails
                    else "Email already exists",
                )
            )
        if not roll:
            problems.append(("roll_number", "Missing Roll Number"))
        elif roll in existing_rolls or roll in seen_rolls:
            problems.append(
                (
                    "roll_number",
                    "Duplicate Roll Number"
                    if roll in seen_rolls
                    else "Roll number already exists",
                )
            )
        if not course:
            problems.append(("course", "Missing Course"))
        if not branch:
            problems.append(("branch", "Missing Branch"))
        if section == "" or section is None:
            problems.append(("section", "Missing Section"))

        clean = {
            "_row": r,
            "first_name": first,
            "last_name": last,
            "email": email,
            "roll_number": roll,
            "course": course,
            "branch": branch,
            "section": section,
            "phone": row.get("phone"),
            "batch": row.get("batch"),
            "academic_year": row.get("academic_year"),
        }

        if problems:
            for field, message in problems:
                errors.append({"row": r, "field": field, "error": message})
        else:
            valid.append(clean)
            seen_emails.add(email)
            seen_rolls.add(roll)

    return valid, errors


class ImportService:
    def __init__(self, db):
        self.db = db
        self.students = StudentService(db)

    async def preview(self, content: bytes) -> dict:
        rows = parse_xlsx(content)
        students_col = self.db.students
        emails = {
            d["email"].lower()
            async for d in _aiter(students_col.find({}, {"email": 1}))
        }
        rolls = {
            d["roll_number"]
            async for d in _aiter(students_col.find({}, {"roll_number": 1}))
        }
        valid, errors = validate_rows(rows, emails, rolls)

        # Store validated payload for the confirm step (short-lived).
        token = await self._store_pending(valid)
        return {
            "import_token": token,
            "total_records": len(rows),
            "valid_records": len(valid),
            "invalid_records": len(errors),
            "duplicate_records": sum(1 for e in errors if "Duplicate" in e["error"]),
            "errors": [
                {"row": e["row"], "field": e["field"], "error": e["error"]}
                for e in errors
            ],
            "sample_valid": valid[:5],
        }

    async def _store_pending(self, valid_rows: list[dict]) -> str:
        doc = {
            "rows": valid_rows,
            "created_at": datetime.now(timezone.utc),
            "expires_at": datetime.fromtimestamp(
                datetime.now(timezone.utc).timestamp() + 900, tz=timezone.utc
            ),
        }
        result = await self.db.import_previews.insert_one(doc)
        return str(result.inserted_id)

    async def confirm(
        self, import_token: str, actor: dict, ip=None, user_agent=None
    ) -> dict:
        try:
            pending = await self.db.import_previews.find_one(
                {"_id": bson.ObjectId(import_token)}
            )
        except Exception:
            raise NotFoundError("Import session not found or expired")
        if not pending:
            raise NotFoundError("Import session not found or expired")
        created = 0
        skipped = 0
        for row in pending.get("rows", []):
            payload = {
                "first_name": row["first_name"],
                "last_name": row["last_name"],
                "email": row["email"],
                "roll_number": row["roll_number"],
                "course": row["course"],
                "branch": row["branch"],
                "section": row["section"],
                "phone": row.get("phone"),
                "batch_id": None,
                "academic_year_id": None,
            }
            batch_name = (row.get("batch") or "").strip()
            year_name = (row.get("academic_year") or "").strip()
            if batch_name or year_name:
                from app.repositories.system import AcademicRepository

                academic = AcademicRepository(self.db)
                year_doc = await academic.ensure_year(year_name or "2025-26")
                payload["academic_year_id"] = str(year_doc["_id"])
                if batch_name:
                    batch_doc = await academic.ensure_batch(
                        batch_name, str(year_doc["_id"])
                    )
                    payload["batch_id"] = str(batch_doc["_id"])
            try:
                await self.students.create_student(
                    payload, actor, ip=ip, user_agent=user_agent
                )
                created += 1
            except Exception:
                skipped += 1

        await self.db.import_previews.delete_one({"_id": pending["_id"]})
        await AuditService(self.db).log(
            AuditAction.BULK_IMPORT,
            user_id=str(actor["_id"]),
            role=actor["role"],
            entity="bulk_import",
            new_data={"created": created, "skipped": skipped},
            ip_address=ip,
            user_agent=user_agent,
        )
        return {"created": created, "skipped": skipped}


async def _aiter(cursor):
    async for item in cursor:
        yield item

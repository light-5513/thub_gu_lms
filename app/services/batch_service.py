"""Batch service for managing student batches and bulk assignments."""

import logging
from io import BytesIO

from bson import ObjectId
from openpyxl import load_workbook
from pymongo import UpdateOne

from app.core.errors import AppError, NotFoundError
from app.repositories.batches import BatchRepository
from app.repositories.users import StudentRepository

logger = logging.getLogger(__name__)

class BatchService:
    def __init__(self, db):
        self.db = db
        self.batches = BatchRepository(db)
        self.students = db.students

    async def list_batches(self, page: int = 1, page_size: int = 50) -> dict:
        """List all batches."""
        total = await self.batches.collection.count_documents({})
        skip = (page - 1) * page_size
        cursor = self.batches.collection.find({}).sort("created_at", -1).skip(skip).limit(page_size)
        items = []
        from datetime import datetime
        async for doc in cursor:
            doc["id"] = str(doc.pop("_id"))
            for k, v in doc.items():
                if isinstance(v, datetime):
                    doc[k] = v.isoformat()
            items.append(doc)
            
        return {
            "items": items,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": (total + page_size - 1) // page_size
        }

    async def create_batch(self, batch_data: dict) -> dict:
        """Create a new batch."""
        return await self.batches.create_batch(batch_data)

    async def get_batch(self, batch_id: str) -> dict:
        """Get a batch by ID."""
        return await self.batches.get(batch_id)

    async def delete_batch(self, batch_id: str) -> None:
        """Delete a batch and unassign all students."""
        await self.batches.delete(batch_id)
        # Unassign students
        await self.students.update_many({"batch_id": batch_id}, {"$set": {"batch_id": None}})

    async def assign_students(self, batch_id: str, roll_numbers: list[str]) -> dict:
        """Assign a list of roll numbers to a batch."""
        # Ensure batch exists
        batch = await self.get_batch(batch_id)
        
        # Find valid roll numbers
        cursor = self.students.find({"roll_number": {"$in": roll_numbers}}, {"roll_number": 1, "batch_id": 1})
        found_students = await cursor.to_list(None)
        found_rolls = {s["roll_number"] for s in found_students}
        
        failed_rolls = list(set(roll_numbers) - found_rolls)
        success_count = len(found_rolls)
        
        if found_rolls:
            await self.students.update_many(
                {"roll_number": {"$in": list(found_rolls)}},
                {"$set": {"batch_id": batch_id}}
            )
            # Update student count (recalculate completely to be safe)
            actual_count = await self.students.count_documents({"batch_id": batch_id})
            await self.batches.update_batch(batch_id, {"student_count": actual_count})
            
        return {
            "success_count": success_count,
            "failed_roll_numbers": failed_rolls
        }

    async def parse_and_assign_xlsx(self, batch_id: str, content: bytes) -> dict:
        """Parse an XLSX file for roll numbers and assign them."""
        wb = load_workbook(filename=BytesIO(content), read_only=True, data_only=True)
        sheet = wb.active
        
        # Assume first row is header, look for "roll_number"
        headers = [str(cell.value).strip().lower() for cell in sheet[1] if cell.value]
        if "roll_number" not in headers:
            raise AppError("Missing 'roll_number' column in Excel file")
            
        roll_index = headers.index("roll_number")
        
        roll_numbers = []
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) > roll_index and row[roll_index]:
                roll_numbers.append(str(row[roll_index]).strip())
                
        if not roll_numbers:
            raise AppError("No roll numbers found in the Excel file")
            
        return await self.assign_students(batch_id, roll_numbers)

    async def get_batch_students(self, batch_id: str, page: int = 1, page_size: int = 50) -> dict:
        """List students in a batch."""
        total = await self.students.count_documents({"batch_id": batch_id})
        skip = (page - 1) * page_size
        cursor = self.students.find({"batch_id": batch_id}).sort("first_name", 1).skip(skip).limit(page_size)
        items = []
        from datetime import datetime
        async for doc in cursor:
            doc["id"] = str(doc.pop("_id"))
            for k, v in doc.items():
                if isinstance(v, datetime):
                    doc[k] = v.isoformat()
            items.append(doc)
            
        return {
            "items": items,
            "total": total,
            "page": page,
            "page_size": page_size,
            "total_pages": (total + page_size - 1) // page_size if page_size else 0
        }

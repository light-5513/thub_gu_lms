"""Report generation: CSV, Excel and PDF exports."""

import csv
import io
from datetime import datetime, timezone

from openpyxl import Workbook


def rows_to_csv(headers: list[str], rows: list[list]) -> bytes:
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(headers)
    writer.writerows(rows)
    return buf.getvalue().encode("utf-8")


def rows_to_xlsx(
    headers: list[str], rows: list[list], sheet_name: str = "Report"
) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    ws.append(headers)
    for row in rows:
        ws.append(row)
    # Auto-ish column widths
    for col_idx, header in enumerate(headers, start=1):
        width = max(len(str(header)), 12)
        ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = min(
            width + 4, 40
        )
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def rows_to_pdf(title: str, headers: list[str], rows: list[list]) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, topMargin=15 * mm, bottomMargin=15 * mm)
    styles = getSampleStyleSheet()
    elements = [Paragraph(title, styles["Title"]), Spacer(1, 6)]
    generated = datetime.now(timezone.utc).strftime("%d %b %Y %H:%M UTC")
    elements.append(Paragraph(f"Generated {generated}", styles["Normal"]))
    elements.append(Spacer(1, 10))

    data = [headers] + [[str(c) for c in row] for row in rows]
    table = Table(data, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1e3a5f")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 8),
                (
                    "ROWBACKGROUNDS",
                    (0, 1),
                    (-1, -1),
                    [colors.white, colors.HexColor("#f2f6fa")],
                ),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    elements.append(table)
    doc.build(elements)
    return buf.getvalue()


# -----------------------------------------------------------------------
# Per-student report (used by /admin/reports/student/{student_id})
# -----------------------------------------------------------------------


def _fmt_dt(value) -> str:
    if value is None:
        return "-"
    if isinstance(value, str):
        return value
    try:
        return value.strftime("%d %b %Y %H:%M")
    except Exception:
        return str(value)


def _fmt_pct(value) -> str:
    return f"{value}%" if value is not None else "-"


def student_report_to_csv(report: dict) -> bytes:
    """Multi-section CSV: Profile, Attendance, Coding, Leaderboard, Recent Records."""
    buf = io.StringIO()
    w = csv.writer(buf)

    profile = report.get("profile", {})
    attendance = report.get("attendance", {})
    coding = report.get("coding", {})
    leaderboard = report.get("leaderboard", {})
    by_subject = report.get("by_subject", [])
    recent = report.get("recent_records", [])

    # --- Profile ---
    w.writerow(["Student Report"])
    w.writerow(["Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
    w.writerow(["Days covered", report.get("days", 90)])
    w.writerow([])
    w.writerow(["Field", "Value"])
    for k, v in [
        ("Name", profile.get("name")),
        ("Roll number", profile.get("roll_number")),
        ("Email", profile.get("email")),
        ("Course", profile.get("course")),
        ("Branch", profile.get("branch")),
        ("Section", profile.get("section")),
        ("Status", profile.get("status")),
        ("Phone", profile.get("phone")),
    ]:
        w.writerow([k, v])
    w.writerow([])

    # --- Attendance summary ---
    w.writerow(["Attendance Summary"])
    w.writerow(["Metric", "Value"])
    for k, v in [
        ("Present", attendance.get("present")),
        ("Late", attendance.get("late")),
        ("Absent", attendance.get("absent")),
        ("Leave", attendance.get("leave")),
        ("Total classes", attendance.get("total_classes")),
        ("Percentage", _fmt_pct(attendance.get("percentage"))),
        ("Threshold", _fmt_pct(attendance.get("threshold"))),
        ("Is low attendance", "Yes" if attendance.get("is_low") else "No"),
    ]:
        w.writerow([k, v])
    w.writerow([])

    # --- Attendance by subject ---
    w.writerow(["Attendance by Subject"])
    w.writerow(["Subject", "Present", "Total", "Percentage"])
    for s in by_subject:
        w.writerow(
            [
                s.get("subject"),
                s.get("present"),
                s.get("total"),
                _fmt_pct(s.get("percentage")),
            ]
        )
    w.writerow([])

    # --- Coding ---
    w.writerow(["Coding Profiles"])
    w.writerow(
        [
            "Platform",
            "Username",
            "Rating",
            "Problems Solved",
            "Sync Status",
            "Last Synced",
        ]
    )
    for p in coding.get("platforms", []):
        w.writerow(
            [
                p.get("platform"),
                p.get("username"),
                p.get("rating") or "-",
                p.get("problems_solved") or 0,
                p.get("sync_status") or "-",
                _fmt_dt(p.get("last_synced") or p.get("fetched_at")),
            ]
        )
    w.writerow([])

    # --- Leaderboard ---
    w.writerow(["Leaderboard"])
    w.writerow(["Metric", "Value"])
    for k, v in [
        ("Overall score", leaderboard.get("overall_score")),
        ("Rank", leaderboard.get("rank")),
        ("Total students", leaderboard.get("total_students")),
        ("Current streak", leaderboard.get("current_streak")),
        ("Longest streak", leaderboard.get("longest_streak")),
    ]:
        w.writerow([k, v])
    w.writerow([])

    # --- Recent attendance records ---
    w.writerow(["Recent Attendance Records"])
    w.writerow(["Date", "Subject", "Status", "Start", "End", "Room"])
    for r in recent:
        w.writerow(
            [
                r.get("date"),
                r.get("subject") or "-",
                r.get("status"),
                r.get("start_time") or "-",
                r.get("end_time") or "-",
                r.get("room") or "-",
            ]
        )

    return buf.getvalue().encode("utf-8")


def student_report_to_xlsx(report: dict) -> bytes:
    """Multi-sheet XLSX: Profile, Attendance, BySubject, Coding, Leaderboard, Recent."""
    wb = Workbook()
    profile = report.get("profile", {})
    attendance = report.get("attendance", {})
    coding = report.get("coding", {})
    leaderboard = report.get("leaderboard", {})
    by_subject = report.get("by_subject", [])
    recent = report.get("recent_records", [])

    # Sheet 1: Profile
    ws = wb.active
    ws.title = "Profile"
    ws.append(["Student Report"])
    ws.append(["Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")])
    ws.append(["Days covered", report.get("days", 90)])
    ws.append([])
    ws.append(["Field", "Value"])
    for k, v in [
        ("Name", profile.get("name")),
        ("Roll number", profile.get("roll_number")),
        ("Email", profile.get("email")),
        ("Course", profile.get("course")),
        ("Branch", profile.get("branch")),
        ("Section", profile.get("section")),
        ("Status", profile.get("status")),
        ("Phone", profile.get("phone")),
    ]:
        ws.append([k, v])

    # Sheet 2: Attendance
    ws2 = wb.create_sheet("Attendance")
    ws2.append(["Metric", "Value"])
    for k, v in [
        ("Present", attendance.get("present")),
        ("Late", attendance.get("late")),
        ("Absent", attendance.get("absent")),
        ("Leave", attendance.get("leave")),
        ("Total classes", attendance.get("total_classes")),
        ("Percentage", attendance.get("percentage")),
        ("Threshold", attendance.get("threshold")),
        ("Is low attendance", "Yes" if attendance.get("is_low") else "No"),
    ]:
        ws2.append([k, v])

    # Sheet 3: By subject
    ws3 = wb.create_sheet("BySubject")
    ws3.append(["Subject", "Present", "Total", "Percentage"])
    for s in by_subject:
        ws3.append(
            [s.get("subject"), s.get("present"), s.get("total"), s.get("percentage")]
        )

    # Sheet 4: Coding
    ws4 = wb.create_sheet("Coding")
    ws4.append(
        [
            "Platform",
            "Username",
            "Rating",
            "Problems Solved",
            "Sync Status",
            "Last Synced",
            "Profile URL",
        ]
    )
    for p in coding.get("platforms", []):
        ws4.append(
            [
                p.get("platform"),
                p.get("username"),
                p.get("rating"),
                p.get("problems_solved"),
                p.get("sync_status"),
                _fmt_dt(p.get("last_synced") or p.get("fetched_at")),
                p.get("profile_url"),
            ]
        )

    # Sheet 5: Leaderboard
    ws5 = wb.create_sheet("Leaderboard")
    ws5.append(["Metric", "Value"])
    for k, v in [
        ("Overall score", leaderboard.get("overall_score")),
        ("Rank", leaderboard.get("rank")),
        ("Total students", leaderboard.get("total_students")),
        ("Current streak", leaderboard.get("current_streak")),
        ("Longest streak", leaderboard.get("longest_streak")),
    ]:
        ws5.append([k, v])

    # Sheet 6: Recent records
    ws6 = wb.create_sheet("RecentRecords")
    ws6.append(["Date", "Subject", "Status", "Start", "End", "Room"])
    for r in recent:
        ws6.append(
            [
                r.get("date"),
                r.get("subject"),
                r.get("status"),
                r.get("start_time"),
                r.get("end_time"),
                r.get("room"),
            ]
        )

    # Reasonable column widths
    for sheet in [ws, ws2, ws3, ws4, ws5, ws6]:
        for col_cells in sheet.columns:
            try:
                letter = col_cells[0].column_letter
                max_len = (
                    max(
                        (len(str(c.value)) if c.value is not None else 0)
                        for c in col_cells
                    )
                    + 3
                )
                sheet.column_dimensions[letter].width = min(max(max_len, 14), 50)
            except Exception:
                pass

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def student_report_to_pdf(report: dict) -> bytes:
    """Single PDF with sections: header, profile, attendance summary, by subject, coding, leaderboard, recent records."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import (
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    profile = report.get("profile", {})
    attendance = report.get("attendance", {})
    coding = report.get("coding", {})
    leaderboard = report.get("leaderboard", {})
    by_subject = report.get("by_subject", [])
    recent = report.get("recent_records", [])

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        topMargin=15 * mm,
        bottomMargin=15 * mm,
        leftMargin=15 * mm,
        rightMargin=15 * mm,
    )
    styles = getSampleStyleSheet()
    title_style = styles["Title"]
    h2 = styles["Heading2"]
    body = styles["BodyText"]

    elements = []
    elements.append(
        Paragraph(f"Student Report — {profile.get('name', '')}", title_style)
    )
    sub = f"Roll No: {profile.get('roll_number') or '—'}  ·  {profile.get('course') or '—'} · {profile.get('branch') or '—'} · Section {profile.get('section') or '—'}"
    elements.append(Paragraph(sub, body))
    elements.append(
        Paragraph(
            f"Generated {datetime.now(timezone.utc).strftime('%d %b %Y %H:%M UTC')}  ·  Last {report.get('days', 90)} days",
            body,
        )
    )
    elements.append(Spacer(1, 6 * mm))

    def _table(data, col_widths=None, header_bg="#1e3a5f"):
        if not data:
            return None
        t = Table(data, colWidths=col_widths, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(header_bg)),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#f6f1e7")],
                    ),
                    ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]
            )
        )
        return t

    # Profile
    elements.append(Paragraph("Profile", h2))
    profile_data = [
        ["Field", "Value"],
        ["Name", profile.get("name", "")],
        ["Email", profile.get("email", "")],
        ["Roll number", profile.get("roll_number", "")],
        [
            "Course / Branch / Section",
            f"{profile.get('course') or ''} / {profile.get('branch') or ''} / {profile.get('section') or ''}",
        ],
        ["Status", profile.get("status", "")],
        ["Phone", profile.get("phone", "") or "—"],
    ]
    elements.append(_table(profile_data, col_widths=[55 * mm, 115 * mm]))
    elements.append(Spacer(1, 4 * mm))

    # Attendance summary
    elements.append(Paragraph("Attendance Summary", h2))
    pct = attendance.get("percentage")
    pct_text = f"{pct}%" if pct is not None else "—"
    is_low = attendance.get("is_low")
    threshold = attendance.get("threshold")
    att_data = [
        [
            "Present",
            "Late",
            "Absent",
            "Leave",
            "Total",
            "Percentage",
            "Threshold",
            "Low?",
        ],
        [
            attendance.get("present", 0),
            attendance.get("late", 0),
            attendance.get("absent", 0),
            attendance.get("leave", 0),
            attendance.get("total_classes", 0),
            pct_text,
            f"{threshold}%" if threshold is not None else "—",
            "Yes" if is_low else "No",
        ],
    ]
    elements.append(_table(att_data, col_widths=[20 * mm] * 8))
    elements.append(Spacer(1, 4 * mm))

    # By subject
    if by_subject:
        elements.append(Paragraph("Attendance by Subject", h2))
        sub_data = [["Subject", "Present", "Total", "%"]]
        for s in by_subject:
            sub_data.append(
                [
                    s.get("subject", ""),
                    s.get("present", 0),
                    s.get("total", 0),
                    f"{s.get('percentage', 0)}%",
                ]
            )
        elements.append(
            _table(sub_data, col_widths=[70 * mm, 25 * mm, 25 * mm, 25 * mm])
        )
        elements.append(Spacer(1, 4 * mm))

    # Coding
    elements.append(Paragraph("Coding Profiles", h2))
    if coding.get("platforms"):
        cdata = [["Platform", "Username", "Rating", "Problems", "Sync", "Last Synced"]]
        for p in coding["platforms"]:
            cdata.append(
                [
                    p.get("platform", ""),
                    p.get("username") or "",
                    str(p.get("rating") or "-"),
                    p.get("problems_solved") or 0,
                    p.get("sync_status") or "-",
                    _fmt_dt(p.get("last_synced") or p.get("fetched_at")),
                ]
            )
        elements.append(
            _table(
                cdata, col_widths=[28 * mm, 50 * mm, 22 * mm, 22 * mm, 20 * mm, 38 * mm]
            )
        )
    else:
        elements.append(Paragraph("No coding profiles connected.", body))
    elements.append(Spacer(1, 4 * mm))

    # Leaderboard
    elements.append(Paragraph("Leaderboard", h2))
    lb_data = [
        ["Overall score", "Rank", "Total students", "Current streak", "Longest streak"],
        [
            leaderboard.get("overall_score", 0),
            leaderboard.get("rank") or "-",
            leaderboard.get("total_students", 0),
            leaderboard.get("current_streak", 0),
            leaderboard.get("longest_streak", 0),
        ],
    ]
    elements.append(
        _table(lb_data, col_widths=[40 * mm, 30 * mm, 40 * mm, 40 * mm, 40 * mm])
    )
    elements.append(Spacer(1, 4 * mm))

    # Recent records
    if recent:
        elements.append(Paragraph("Recent Attendance Records", h2))
        rec_data = [["Date", "Subject", "Status", "Start", "End", "Room"]]
        for r in recent:
            rec_data.append(
                [
                    r.get("date", ""),
                    r.get("subject") or "-",
                    r.get("status", ""),
                    r.get("start_time") or "-",
                    r.get("end_time") or "-",
                    r.get("room") or "-",
                ]
            )
        elements.append(
            _table(
                rec_data,
                col_widths=[28 * mm, 50 * mm, 25 * mm, 20 * mm, 20 * mm, 27 * mm],
            )
        )

    doc.build(elements)
    return buf.getvalue()
